#!/usr/bin/env python3
"""Recalculate an XLSX file with an existing LibreOffice installation."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile

from openpyxl import load_workbook


ERROR_VALUES = {
    "#VALUE!",
    "#DIV/0!",
    "#REF!",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#N/A",
}

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script"
 script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>
"""


def scan_workbook(path: Path) -> dict[str, object]:
    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    formula_count = 0
    errors: dict[str, list[str]] = {value: [] for value in ERROR_VALUES}

    try:
        for formula_sheet, value_sheet in zip(
            formulas.worksheets, values.worksheets, strict=True
        ):
            for formula_row, value_row in zip(
                formula_sheet.iter_rows(), value_sheet.iter_rows(), strict=True
            ):
                for formula_cell, value_cell in zip(
                    formula_row, value_row, strict=True
                ):
                    if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                        formula_count += 1
                    if isinstance(value_cell.value, str) and value_cell.value in ERROR_VALUES:
                        errors[value_cell.value].append(
                            f"{value_sheet.title}!{value_cell.coordinate}"
                        )
    finally:
        formulas.close()
        values.close()

    summary = {
        key: {"count": len(locations), "locations": locations[:20]}
        for key, locations in sorted(errors.items())
        if locations
    }
    total_errors = sum(item["count"] for item in summary.values())
    return {
        "status": "success" if total_errors == 0 else "errors_found",
        "total_formulas": formula_count,
        "total_errors": total_errors,
        "error_summary": summary,
    }


def find_soffice() -> str | None:
    candidates = [
        os.environ.get("SOFFICE"),
        "/usr/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        shutil.which("soffice"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file() and os.access(candidate, os.X_OK):
            return candidate
    return None


def recalculate(path: Path, timeout: int) -> dict[str, object]:
    if path.suffix.lower() != ".xlsx":
        return {"status": "unsupported", "error": "recalculation supports XLSX files only"}
    if not path.is_file():
        return {"status": "missing", "error": f"file does not exist: {path}"}

    soffice = find_soffice()
    if not soffice:
        return {"status": "unavailable", "error": "soffice is not available"}

    with tempfile.TemporaryDirectory(prefix="xlsx-recalc-") as temporary:
        root = Path(temporary)
        profile = root / "profile"
        working = root / path.name
        shutil.copy2(path, working)

        profile_uri = profile.as_uri()
        initialize = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--terminate_after_init",
            ],
            capture_output=True,
            text=True,
            timeout=timeout,
            check=False,
        )
        if initialize.returncode != 0:
            return {
                "status": "failed",
                "error": initialize.stderr.strip() or "LibreOffice profile initialization failed",
            }

        macro_dir = profile / "user" / "basic" / "Standard"
        macro_dir.mkdir(parents=True, exist_ok=True)
        (macro_dir / "Module1.xba").write_text(MACRO, encoding="utf-8")

        command = [
            soffice,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--norestore",
            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave"
            "?language=Basic&location=application",
            str(working),
        ]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=timeout,
                check=False,
            )
        except subprocess.TimeoutExpired:
            return {"status": "failed", "error": f"recalculation exceeded {timeout} seconds"}

        if result.returncode != 0:
            return {
                "status": "failed",
                "error": result.stderr.strip() or "LibreOffice recalculation failed",
            }

        validation = scan_workbook(working)
        if validation["status"] == "success":
            shutil.copy2(working, path)
        else:
            validation["error"] = "formula errors remain; original file was not replaced"
        return validation


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--timeout", type=int, default=30)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        result = recalculate(args.file.resolve(), args.timeout)
    except Exception as error:
        result = {"status": "failed", "error": str(error)}
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0 if result.get("status") == "success" else 1


if __name__ == "__main__":
    raise SystemExit(main())
